"""Copy of the deck for a static PDF preview: everything the teacher reveals by clicking through
stays visible, while trigger feedback (answers, overlays, pop-ups) is left out."""
import sys
from pptx import Presentation
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
prs = Presentation(sys.argv[1])
for slide in prs.slides:
    first = {}
    for ctn in slide._element.iter(P + "cTn"):
        cls = ctn.get("presetClass")
        if not cls:
            continue
        inter = False
        anc = ctn.getparent()
        while anc is not None:
            if anc.tag == P + "cTn" and anc.get("nodeType") == "interactiveSeq":
                inter = True
                break
            anc = anc.getparent()
        for tgt in ctn.iter(P + "spTgt"):
            first.setdefault(tgt.get("spid"), (cls, inter))
            break
    for shp in list(slide.shapes):
        cls_inter = first.get(str(shp.shape_id))
        if cls_inter and cls_inter[0] == "entr" and cls_inter[1]:
            shp._element.getparent().remove(shp._element)
prs.save(sys.argv[2])
