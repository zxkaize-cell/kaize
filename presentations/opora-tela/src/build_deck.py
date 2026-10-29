"""Builds «Опора тела и движение» — interactive lesson deck (python-pptx + hand-written timing XML).

Design: pictures, designer backgrounds and typography only (Unbounded / Nunito / Caveat) —
no icons or drawn illustrations.
"""
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

import kit
from kit import (C, E, FONT, NUNITO, DISPLAY, HAND, ASSETS, IMG, text, shape_text, shadow, bevel, transparent_fill,
                 set_name, set_descr, link_slide, link_url, notes, rgb, text_width_in, glass, badge, label)
from pptx_fx import Timeline, Fx, set_transition

OUT = sys.argv[1] if len(sys.argv) > 1 else "deck.pptx"
GAME_URL = os.environ.get("GAME_URL", "https://kostik-skeleton.lovable.app")
VIDEO_URL = "https://youtu.be/6C_RWeZtCNg?si=fnJgUxPPG3wl2pDA"
SW, SH = 13.333, 7.5
SOUNDS = {k: os.path.join(ASSETS, k + ".wav") for k in ("correct", "wrong", "win", "flip", "pop")}
COMP = os.path.join(IMG, "comp")
INK = C["ink"]
MINT_DARK = "13846F"
MINT = "13A386"


def img(key):
    for ext in (".png", ".jpg"):
        p = os.path.join(IMG, key + ext)
        if os.path.exists(p):
            return p
    raise FileNotFoundError(key)


def comp(name):
    return os.path.join(COMP, name + ".jpg")


prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
cp = prs.core_properties
cp.title = "Опора тела и движение"
cp.author = "Шелешко А.И."
cp.subject = "Окружающий мир, 3 класс — интерактивный урок"
cp.keywords = "скелет, мышцы, суставы, осанка"
LAYOUT_BLANK = prs.slide_layouts[6]
N = 31
S = [None] + [prs.slides.add_slide(LAYOUT_BLANK) for _ in range(N)]
HOME, HUB = 5, 7


# ---------------------------------------------------------------- building blocks
def remove_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def shp(slide, kind, x, y, w, h, fill=None, name=None, radius=None, line=None, line_w=0.0):
    t = {"round": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL, "rect": MSO_SHAPE.RECTANGLE}[kind]
    s = slide.shapes.add_shape(t, E(x), E(y), E(w), E(h))
    remove_style(s)
    if kind == "round":
        s.adjustments[0] = min(0.5, (radius if radius is not None else min(w, h) / 2) / min(w, h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = E(line_w)
    else:
        s.line.fill.background()
    if name:
        set_name(s, name)
    return s


def bg(slide, name):
    p = slide.shapes.add_picture(os.path.join(ASSETS, name), 0, 0, prs.slide_width, prs.slide_height)
    set_name(p, "!!bg " + name.split(".")[0])
    set_descr(p, "")
    tree = slide.shapes._spTree
    tree.remove(p._element)
    tree.insert(2, p._element)
    return p


def title(slide, txt, color=INK, size=28, x=0.6, y=0.72, w=10.3, h=1.0, align="l"):
    return text(slide, x, y, w, h, txt, size, color, DISPLAY, False, align, "t", name="Title", spacing=1.05)


def kicker(slide, txt, color, x=0.6, y=0.4):
    t = text(slide, x, y, 9.0, 0.3, txt, 12, color, FONT, True, name="kicker", check=False)
    for r in t._element.iter(qn("a:rPr")):
        r.set("spc", "200")
    return t


def pill_btn(slide, caption, x, y, w, h, fill, color, target=None, url=None, name=None, size=12, face=DISPLAY,
             glass_btn=False):
    if glass_btn:
        b = label(slide, x, y, w, h, caption, size, color, face, fill="FFFFFF", alpha=82, gloss=False,
                  margins=(0.08, 0.0), name=name or caption, check=False)
    else:
        b = label(slide, x, y, w, h, caption, size, color, face, fill=fill, margins=(0.08, 0.0), name=name or caption,
                  check=False)
    if target:
        link_slide(b, S[target])
    if url:
        link_url(b, url)
    return b


def nav(slide, museum=False):
    pill_btn(slide, "МАРШРУТ", 11.35, 0.36, 1.4, 0.44, None, INK, target=HOME, name="!!nav route", size=11, glass_btn=True)
    if museum:
        pill_btn(slide, "МУЗЕЙ", 9.95, 0.36, 1.25, 0.44, None, INK, target=HUB, name="!!nav museum", size=11, glass_btn=True)


def next_btn(slide, caption, target, y=6.72):
    w = text_width_in(caption + "  »", 12, DISPLAY) + 0.6
    return pill_btn(slide, caption + "  »", 12.73 - w, y, w, 0.52, C["sun"], INK, target=target, name="!!next")


def photo(slide, key_or_path, x, y, w, h, name, focus=(0.5, 0.5), radius=0.22, border="FFFFFF", border_w=0.05,
          descr="", sh=True):
    path = key_or_path if os.path.sep in key_or_path else img(key_or_path)
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
    target = w / h
    if iw / ih > target:
        vis = target / (iw / ih)
        left = min(max(focus[0] - vis / 2, 0), 1 - vis)
        pic.crop_left, pic.crop_right = left, 1 - vis - left
    else:
        vis = (iw / ih) / target
        top = min(max(focus[1] - vis / 2, 0), 1 - vis)
        pic.crop_top, pic.crop_bottom = top, 1 - vis - top
    geom = pic._element.spPr.find(qn("a:prstGeom"))
    geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = geom.makeelement(qn("a:avLst"), {})
        geom.append(av)
    for g in list(av):
        av.remove(g)
    av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": "val %d" % int(min(50000, radius / min(w, h) * 100000))}))
    if border:
        pic.line.color.rgb = rgb(border)
        pic.line.width = E(border_w)
    if sh:
        shadow(pic, blur=0.35, dist=0.1, alpha=22, color="2B2A5A")
    set_name(pic, name)
    set_descr(pic, descr)
    return pic


def mascot(slide, key, cx, bottom, h, name=None):
    p = img(key)
    w_px, h_px = Image.open(p).size
    w = h * w_px / h_px
    pic = slide.shapes.add_picture(p, E(cx - w / 2), E(bottom - h), height=E(h))
    set_name(pic, name or key)
    set_descr(pic, "Костик — весёлый скелетик-экскурсовод")
    return pic


def hit(slide, x, y, w, h, name, target=None, url=None):
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
    remove_style(r)
    transparent_fill(r)
    set_name(r, name)
    if target:
        link_slide(r, S[target])
    if url:
        link_url(r, url)
    return r


def num(slide, t, x, y, w, h, color, size=30, align="l", name="number", anchor="t"):
    return text(slide, x, y, w, h, t, size, color, DISPLAY, False, align=align, anchor=anchor, name=name, check=False)


def fx(shape, kind, dur=500, delay=0, sound=None, **kw):
    return Fx(shape, kind, dur, delay, sound, **kw)


def timeline(slide):
    return Timeline(slide, SOUNDS)


def big_title(slide, lines, y=1.7):
    return text(slide, 0.8, y, 7.8, 2.4, [[(line, {"color": col})] for line, col in lines], 56, "FFFFFF", DISPLAY,
                name="Title", spacing=0.98, extrude={"depth": 0.05, "color": "140F3C"},
                glow_shadow={"dy": 0.08, "blur": 0.08, "alpha": 120, "color": "0B0826"})


SKEL = img("skeleton")
ARM = img("arm")
MANN = img("mannequin")
REC = img("records")

# ================================================================ SLIDE 1 — title
s = S[1]
b1 = bg(s, "bg2_hero.jpg")
k = kicker(s, "ОКРУЖАЮЩИЙ МИР  ·  3 КЛАСС", C["sun"], x=0.8, y=1.25)
ttl = big_title(s, [("Опора тела", "FFFFFF"), ("и движение", C["sun"])])
sub = text(s, 0.8, 4.2, 7.0, 1.0, "Путешествие в музей «Тело человека» вместе с Костиком", 30, "E9E4FF", HAND,
           name="subtitle", spacing=0.95)
auth = glass(s, 0.8, 5.45, 4.7, 1.05, alpha=14, radius=0.3, name="author card", sh=False)
at = text(s, 1.1, 5.45, 4.2, 1.05, [[("Автор: Шелешко А.И.", {"size": 19, "bold": True, "color": "FFFFFF"})],
                                    [("группа П-31", {"size": 15, "color": "D9D2FF", "bold": True})]],
          name="author text", anchor="m")
m1 = mascot(s, "mascot_wave", 10.45, 7.25, 6.1, name="!!mascot")
tl = timeline(s)
tl.begin([fx(b1, "grow", 22000, scale=108000), fx(k, "fade", 600), fx(ttl, "pop", 900, delay=150),
          fx(m1, "bounce", 1300, delay=350)])
tl.then([fx(sub, "float_up", 700)])
tl.then([fx(auth, "float_up", 500), fx(at, "float_up", 500)])
tl.then([fx(m1, "bob", 1700, dy="-0.018")])
tl.apply()
set_transition(s, "curtains", 2200)
notes(s, "Здравствуйте, ребята! Сегодня у нас необычный урок: мы отправимся в путешествие в музей «Тело человека». "
         "А экскурсоводом будет весёлый скелетик Костик. Посмотрите, как он нам машет! Помашем ему в ответ.\n"
         "Тему урока пока не называем — дети отгадают её сами на следующих слайдах.")

# ================================================================ SLIDE 2 — riddle
s = S[2]
bg(s, "bg2_lilac.jpg")
kicker(s, "РАЗМИНКА ДЛЯ УМА", C["violet"])
title(s, "Отгадайте загадку")
cd = glass(s, 0.6, 1.85, 7.9, 4.95, name="riddle card")
qm = text(s, 0.85, 1.55, 1.5, 1.2, "«", 88, C["lilac"], DISPLAY, name="quote mark", check=False)
rid = text(s, 1.3, 2.85, 6.9, 2.6, "Он внутри у нас таится,\nПомогает не свалиться,\nДержит тело, как каркас.\nЧто же это есть у нас?",
           38, INK, HAND, spacing=0.95, name="riddle")
ans = label(s, 1.25, 5.6, 3.3, 0.85, "СКЕЛЕТ", 26, INK, DISPLAY, fill=C["sun"], name="answer")
ab = badge(s, SKEL, 4.85, 5.25, 1.45, center=(0.5, 0.33), side=0.62, name="answer photo", descr="Скелет человека")
m2 = mascot(s, "mascot_think", 10.75, 7.25, 5.7, name="!!mascot")
tl = timeline(s)
tl.begin([fx(cd, "float_up", 600), fx(qm, "fade", 600, delay=200), fx(rid, "fade", 900, delay=350),
          fx(m2, "fly_right", 800, delay=200)])
tl.click([fx(ans, "pop", 700, sound="correct"), fx(ab, "spin_in", 900, delay=200)])
tl.apply()
set_transition(s, "gallery", 1600, dir="l")
notes(s, "Костик приготовил для вас загадку. Послушайте внимательно (учитель читает загадку).\n"
         "Ответ детей: скелет. По щелчку появляется ответ «СКЕЛЕТ» и фотография скелета.\n"
         "— Правильно! Скелет — это каркас нашего тела. А зачем он нам нужен? Сейчас разберёмся.")

# ================================================================ SLIDE 3 — problem question
s = S[3]
bg(s, "bg2_sky.jpg")
kicker(s, "ПРОБЛЕМНЫЙ ВОПРОС", C["coral"])
title(s, "Почему медуза растекается, а мальчик прыгает?", size=26, w=10.6, h=1.1)
nav(s)
pj = photo(s, "jelly", 0.6, 2.0, 7.3, 4.25, "jelly photo", focus=(0.55, 0.5),
           descr="Медуза растеклась на песке, а мальчик высоко прыгает")
ca = glass(s, 8.25, 2.0, 4.48, 2.0, name="card jelly")
ba = badge(s, img("jelly"), 8.45, 2.3, 1.35, center=(0.42, 0.86), side=0.42, name="badge jelly")
ta = text(s, 9.95, 2.0, 2.65, 2.0, "У медузы нет скелета — на суше её мягкое тело не держит форму", 15, INK, FONT, True,
          anchor="m", name="text jelly")
cb = glass(s, 8.25, 4.25, 4.48, 2.0, name="card human")
bb = badge(s, img("jelly"), 8.45, 4.55, 1.35, center=(0.73, 0.4), side=0.78, name="badge boy")
tb = text(s, 9.95, 4.25, 2.65, 2.0, [[("У человека есть ", {}), ("скелет", {"color": C["violet"]}), (" и ", {}),
                                      ("мышцы", {"color": C["coral"]}), (" — это опора тела и движение!", {})]],
          15, INK, FONT, True, name="text human", anchor="m")
kit.check_fit("problem B", "У человека есть скелет и мышцы — это опора тела и движение!", 15, 2.65, 2.0, FONT, True)
tl = timeline(s)
tl.begin([fx(pj, "zoom", 700)])
tl.click([fx(ca, "float_up", 500), fx(ba, "pop", 600, delay=150), fx(ta, "float_up", 500)])
tl.click([fx(cb, "float_up", 500, sound="pop"), fx(bb, "pop", 600, delay=150), fx(tb, "float_up", 500)])
tl.apply()
set_transition(s, "ripple", 1400, dir="center")
notes(s, "Посмотрите на картинку. Медуза в море плавает красиво, а на берегу растекается, как желе. "
         "А мальчик прыгает, бегает и стоит ровно. Почему так?\n"
         "Выслушать версии детей. 1-й щелчок: у медузы нет скелета. 2-й щелчок: у человека есть скелет и мышцы.\n"
         "— Значит, тема нашего урока — «Опора тела и движение».")

# ================================================================ SLIDE 4 — topic, goal, tasks
s = S[4]
bg(s, "bg2_mint.jpg")
kicker(s, "ТЕМА УРОКА", MINT_DARK)
title(s, "Опора тела и движение")
nav(s)
gc = photo(s, comp("goal_card"), 0.6, 1.85, 4.55, 4.85, "goal card", radius=0.3, descr="Дети занимаются спортом")
gh = text(s, 0.95, 4.05, 3.9, 0.55, "Цель урока", 22, C["sun"], DISPLAY, name="goal head")
gt = text(s, 0.95, 4.65, 3.9, 1.85, "Узнать, что помогает нашему телу держать форму и двигаться", 20, "FFFFFF", FONT, True,
          spacing=1.05, name="goal text")
th = text(s, 5.55, 1.85, 7.0, 0.5, "Задачи", 22, INK, DISPLAY, name="tasks head")
tasks = [("Познакомиться со скелетом и его частями", C["sky"]), ("Узнать, как работают суставы и мышцы", C["coral"]),
         ("Выяснить, что такое осанка и как её беречь", MINT), ("Проверить свои знания в играх", "E39B00")]
tl = timeline(s)
tl.begin([fx(gc, "zoom", 700), fx(gh, "fade", 500, delay=350), fx(gt, "fade", 600, delay=450),
          fx(th, "fade", 500, delay=500)])
for i, (t, col) in enumerate(tasks):
    y = 2.5 + i * 1.07
    c = glass(s, 5.55, y, 7.18, 0.92, name="task %d" % (i + 1))
    n = num(s, "0%d" % (i + 1), 5.8, y, 1.0, 0.92, col, 24, anchor="m", name="task num %d" % (i + 1))
    tt = text(s, 6.85, y, 5.75, 0.92, t, 18, INK, FONT, True, anchor="m", name="task text %d" % (i + 1))
    tl.click([fx(c, "fly_right", 500), fx(n, "fly_right", 500), fx(tt, "fly_right", 500)])
tl.apply()
set_transition(s, "prism", 1500, isContent="1")
notes(s, "Тема урока: «Опора тела и движение».\nЦель: узнать, что помогает нашему телу держать форму и двигаться.\n"
         "Задачи появляются по щелчку: 1) познакомиться со скелетом и его частями; 2) узнать, как работают суставы и мышцы; "
         "3) выяснить, что такое осанка и как её беречь; 4) проверить свои знания в играх.\n"
         "Для учителя (цели урока): образовательная — сформировать представление об опорно-двигательной системе человека "
         "(скелет, суставы, мышцы, осанка); развивающая — развивать наблюдательность, умение сравнивать и делать выводы; "
         "воспитательная — воспитывать бережное отношение к своему здоровью.")

# ================================================================ SLIDE 5 — contents (route of the lesson)
s = S[5]
bg(s, "bg2_sky.jpg")
kicker(s, "СОДЕРЖАНИЕ", C["violet"])
title(s, "Маршрут нашего урока")
hint = text(s, 0.6, 1.3, 10.0, 0.55, "Нажми на картинку — и мы сразу туда отправимся!", 25, C["violet"], HAND,
            name="route hint")
STATIONS = [("cinema", (0.5, 0.42), "Видеоэкскурсия", 6),
            ("museum", (0.5, 0.5), "Виртуальная экскурсия", 7),
            (os.path.join(ASSETS, "fizminutka_poster.jpg"), (0.55, 0.4), "Физминутка", 15),
            (comp("tile_games"), (0.5, 0.5), "Игры «Проверь себя»", 16),
            (comp("tile_online"), (0.5, 0.5), "Онлайн-задание", 24),
            ("sport", (0.62, 0.5), "Итоги урока", 25),
            (comp("tile_reflect"), (0.5, 0.5), "Рефлексия", 28)]
TW, TH, PH_ = 2.86, 2.32, 1.62
STATION_COLS = [C["coral"], C["sky"], MINT, C["violet"], C["orange"], "E39B00", C["lilac"]]
pos = [(0.6 + i * 3.09, 1.95) for i in range(4)] + [(0.6 + i * 3.09, 4.55) for i in range(3)]
tl = timeline(s)
step = []
for i, (src, focus, cap, target) in enumerate(STATIONS):
    x, y = pos[i]
    card = glass(s, x, y, TW, TH, name="station card %d" % (i + 1))
    ph_ = photo(s, src, x + 0.1, y + 0.1, TW - 0.2, PH_, "station photo %d" % (i + 1), focus=focus, radius=0.18,
                border=None, sh=False)
    nn = num(s, "0%d" % (i + 1), x + 0.2, y + 0.1 + PH_, 0.65, TH - PH_ - 0.1, STATION_COLS[i], 16, anchor="m",
             name="station num %d" % (i + 1))
    lb = text(s, x + 0.82, y + 0.1 + PH_, TW - 0.95, TH - PH_ - 0.1, cap, 14, INK, FONT, True, anchor="m",
              name="station label %d" % (i + 1), spacing=0.95)
    hit(s, x, y, TW, TH, "station link %d" % (i + 1), target=target)
    d = i * 170
    step += [fx(card, "pop", 600, delay=d), fx(ph_, "pop", 600, delay=d), fx(nn, "fade", 400, delay=d + 300),
             fx(lb, "fade", 400, delay=d + 250)]
tl.begin(step)
m5 = mascot(s, "mascot_point", 11.35, 7.3, 2.95, name="!!mascot")
tl.then([fx(m5, "bounce", 900)])
tl.apply()
set_transition(s, "flythrough", 1300)
notes(s, "Это маршрут нашего урока (интерактивное содержание). Каждая картинка — кнопка: нажмите на неё, "
         "и презентация перейдёт к этой части урока.\n"
         "На всех следующих слайдах есть кнопка «МАРШРУТ» — она возвращает сюда.\n"
         "01 — Видеоэкскурсия, 02 — Виртуальная экскурсия по музею, 03 — Физминутка, 04 — Игры «Проверь себя», "
         "05 — Онлайн-задание, 06 — Итоги урока, 07 — Рефлексия.")

# ================================================================ SLIDE 6 — video excursion
s = S[6]
bg(s, "bg2_cinema.jpg")
kicker(s, "ВИДЕОЭКСКУРСИЯ", C["sun"])
title(s, "Смотрим и запоминаем!", color="FFFFFF")
nav(s)
cin = photo(s, "cinema", 0.6, 1.85, 7.6, 4.3, "cinema", focus=(0.5, 0.45), descr="Уютный кинозал с большим экраном")
play = label(s, 2.55, 3.05, 3.7, 0.8, "СМОТРЕТЬ ВИДЕО  »", 15, "FFFFFF", DISPLAY, fill=C["coral"], name="play button")
hit(s, 0.6, 1.85, 7.6, 4.3, "video link", url=VIDEO_URL)
qc = glass(s, 8.55, 1.85, 4.18, 3.1, alpha=94, name="questions card")
qh = text(s, 8.85, 2.05, 3.7, 0.5, "Найди в фильме ответы:", 15, INK, DISPLAY, name="questions head")
qs = text(s, 8.85, 2.65, 3.7, 2.2, "1. Что служит опорой тела?\n2. Какие кости защищают мозг, сердце и лёгкие?\n3. Что помогает нам двигаться?",
          16, INK, FONT, True, para_space=6, name="questions")
qr_bg = glass(s, 8.55, 5.15, 1.6, 1.6, alpha=100, radius=0.18, name="qr card")
qr = s.shapes.add_picture(os.path.join(ASSETS, "qr_video.png"), E(8.65), E(5.25), E(1.4), E(1.4))
set_name(qr, "qr video")
set_descr(qr, "QR-код на видеоэкскурсию")
link_url(qr, VIDEO_URL)
qt = text(s, 10.35, 5.15, 2.4, 1.6, "или отсканируй QR-код телефоном", 26, "FFFFFF", HAND, anchor="m", name="qr text",
          spacing=0.9)
m6 = mascot(s, "mascot_popcorn", 1.35, 7.35, 2.45, name="!!mascot")
tl = timeline(s)
tl.begin([fx(cin, "zoom", 700), fx(m6, "fly_left", 700, delay=300)])
tl.then([fx(play, "pop", 600, sound="pop")])
tl.then([fx(qc, "float_up", 500), fx(qh, "float_up", 500), fx(qs, "float_up", 500), fx(qr_bg, "float_up", 500, delay=200),
         fx(qr, "float_up", 500, delay=200), fx(qt, "float_up", 500, delay=200)])
tl.then([fx(play, "pulse", 400, scale=107000)])
tl.apply()
set_transition(s, "warp", 1400, dir="in")
notes(s, "Перед просмотром дайте детям задание: найти в фильме ответы на три вопроса (они на слайде).\n"
         "Видео открывается по нажатию на экран или кнопку «Смотреть видео» (нужен интернет): " + VIDEO_URL + "\n"
         "После просмотра обсудите ответы: опора тела — скелет; мозг защищает череп, сердце и лёгкие — грудная клетка; "
         "двигаться помогают мышцы и суставы.")

# ================================================================ SLIDE 7 — museum hub
HALLS = [("Скелет — опора тела", C["sky"], 8), ("Из чего состоит скелет", C["violet"], 9), ("Суставы", MINT, 10),
         ("Мышцы", C["coral"], 11), ("Осанка", C["orange"], 12), ("Музей рекордов", "E39B00", 14)]
s = S[7]
bg(s, "bg2_sky.jpg")
kicker(s, "ВИРТУАЛЬНАЯ ЭКСКУРСИЯ", C["sky"])
title(s, "Музей «Тело человека»")
nav(s)
mu = photo(s, "museum", 0.6, 1.85, 6.35, 3.62, "museum", focus=(0.5, 0.5), descr="Яркий детский музей о теле человека")
wc = glass(s, 0.6, 5.65, 6.35, 1.2, name="welcome card")
wt = text(s, 0.9, 5.65, 4.0, 1.2, "Я Костик — ваш экскурсовод. Выбирайте зал — и вперёд!", 25, INK, HAND, anchor="m",
          name="welcome text", spacing=0.9)
m7 = mascot(s, "mascot_point", 6.05, 7.3, 3.7, name="!!mascot")
tl = timeline(s)
tl.begin([fx(mu, "zoom", 700), fx(wc, "float_up", 500, delay=300), fx(wt, "float_up", 500, delay=300),
          fx(m7, "bounce", 1000, delay=200)])
step = []
for i, (name, col, target) in enumerate(HALLS):
    cx = 7.3 + (i % 2) * 2.77
    cy = 1.85 + (i // 2) * 1.7
    t = photo(s, comp("ticket%d" % (i + 1)), cx, cy, 2.63, 1.52, "!!hall%d" % (i + 1), radius=0.2, border="FFFFFF",
              border_w=0.04, descr="Зал %d: %s" % (i + 1, name))
    zl = text(s, cx + 0.18, cy + 0.72, 2.3, 0.3, "ЗАЛ %d" % (i + 1), 11, "FFFFFF", DISPLAY, name="hall%d num" % (i + 1),
              check=False)
    for r in zl._element.iter(qn("a:rPr")):
        r.set("spc", "150")
    nm = text(s, cx + 0.18, cy + 0.98, 2.35, 0.5, name, 14, "FFFFFF", FONT, True, name="hall%d name" % (i + 1),
              spacing=0.95, check=False)
    kit.check_fit("hall name", name, 14, 2.35, 0.5, FONT, True, spacing=0.95)
    hit(s, cx, cy, 2.63, 1.52, "hall%d link" % (i + 1), target=target)
    d = 200 + i * 150
    step += [fx(t, "pop", 600, delay=d), fx(zl, "fade", 400, delay=d + 250), fx(nm, "fade", 400, delay=d + 250)]
tl.then(step)
tl.apply()
set_transition(s, "doors", 1600, dir="vert")
notes(s, "Мы входим в музей «Тело человека». Здесь шесть залов. Нажмите на билет зала — и Костик проведёт вас туда "
         "(картинка билета плавно «раскрывается» в зал). Из каждого зала можно вернуться сюда кнопкой «МУЗЕЙ».\n"
         "Удобный порядок: залы с 1-го по 6-й, кнопка «Следующий зал» ведёт дальше по порядку.")


# ================================================================ hall frame
def hall_frame(slide, n, ttl_text, bgname, title_size=28):
    name, col, target = HALLS[n - 1]
    bg(slide, bgname)
    kicker(slide, "ЗАЛ %d  ·  ВИРТУАЛЬНАЯ ЭКСКУРСИЯ" % n, col)
    title(slide, ttl_text, size=title_size)
    nav(slide, museum=True)


# ================================================================ SLIDE 8 — hall 1
s = S[8]
hall_frame(s, 1, "Скелет — опора тела", "bg2_sky.jpg")
tp = photo(s, "tent", 0.6, 1.85, 6.6, 3.85, "!!hall1", focus=(0.55, 0.55),
           descr="Палатка без каркаса упала, а палатка с каркасом стоит")
tc = text(s, 0.6, 5.82, 6.6, 0.85, "Без каркаса палатка падает. Без скелета мы не смогли бы даже стоять!", 25,
          C["violet"], HAND, name="tent caption", spacing=0.9)
dc = glass(s, 7.55, 1.85, 5.18, 2.35, name="definition card")
dt = text(s, 7.85, 1.95, 4.65, 2.15, [[("Скелет", {"color": C["violet"]}),
                                       (" — это прочный каркас тела из костей. Он держит форму тела, защищает внутренние органы "
                                        "и вместе с мышцами помогает двигаться.", {})]], 17, INK, FONT, True, name="definition",
          anchor="m")
kit.check_fit("def", "Скелет — это прочный каркас тела из костей. Он держит форму тела, защищает внутренние органы и вместе с мышцами помогает двигаться.", 17, 4.65, 2.15, FONT, True)
st1 = glass(s, 7.55, 4.4, 2.49, 2.1, name="stat adult")
n1 = num(s, "206", 7.6, 4.55, 2.39, 0.9, C["violet"], 40, align="c", name="stat adult num")
l1 = text(s, 7.75, 5.45, 2.1, 0.95, "костей у взрослого человека", 15, INK, FONT, True, align="c", name="stat adult label")
st2 = glass(s, 10.24, 4.4, 2.49, 2.1, name="stat baby")
n2 = num(s, "≈300", 10.29, 4.55, 2.39, 0.9, C["coral"], 40, align="c", name="stat baby num")
l2 = text(s, 10.39, 5.45, 2.2, 1.0, "у малыша — потом часть костей срастается", 14, INK, FONT, True, align="c",
          name="stat baby label")
next_btn(s, "Следующий зал", 9)
tl = timeline(s)
tl.begin([fx(tc, "fade", 600, delay=300)])
tl.click([fx(dc, "float_up", 500), fx(dt, "float_up", 500)])
tl.click([fx(st1, "pop", 600, sound="pop"), fx(n1, "pop", 600), fx(l1, "fade", 400, delay=200),
          fx(st2, "pop", 600, delay=250), fx(n2, "pop", 600, delay=250), fx(l2, "fade", 400, delay=450)])
tl.apply()
set_transition(s, "morph", 1600)
notes(s, "Зал 1. Посмотрите на палатки: у одной нет каркаса — и она упала. Так и наше тело: без скелета мы не смогли бы стоять.\n"
         "Скелет — это прочный каркас тела из костей. Он держит форму тела, защищает внутренние органы и вместе с мышцами помогает двигаться.\n"
         "Щелчок — определение, ещё щелчок — интересные числа: у взрослого 206 костей, а у новорождённого около 300 — "
         "по мере роста некоторые кости срастаются.")

# ================================================================ SLIDE 9 — hall 2 (hotspots)
s = S[9]
hall_frame(s, 2, "Из чего состоит скелет", "bg2_lilac.jpg")
hn = text(s, 0.6, 1.6, 8.0, 0.5, "Нажимай на названия — узнаешь секреты костей!", 24, C["violet"], HAND, name="hotspot hint")
SKH = 5.0
SKW = SKH * 3 / 4
SKX, SKY = (SW - SKW) / 2, 2.15
sk = photo(s, "skeleton", SKX, SKY, SKW, SKH, "!!hall2", descr="Скелет человека спереди")
PTS = {"skull": (0.455, 0.11), "ribs": (0.425, 0.31), "hand": (0.215, 0.56), "spine": (0.505, 0.43),
       "pelvis": (0.575, 0.5), "legs": (0.585, 0.67)}
PARTS = [  # side, row, key, label, colour, text
    ("l", 0, "skull", "Череп", C["violet"], "Защищает головной мозг — как прочный шлем"),
    ("l", 1, "ribs", "Грудная клетка", C["coral"], "12 пар рёбер берегут сердце и лёгкие"),
    ("l", 2, "hand", "Кости рук", C["sky"], "Берём, пишем, рисуем. В одной кисти 27 косточек!"),
    ("r", 0, "spine", "Позвоночник", MINT, "Главная опора тела: 33–34 позвонка, поэтому он гибкий"),
    ("r", 1, "pelvis", "Кости таза", C["orange"], "Опора для позвоночника, к ним крепятся ноги"),
    ("r", 2, "legs", "Кости ног", C["violet"], "Держат вес тела. Бедренная кость — самая длинная!"),
]
tl = timeline(s)
tl.begin([fx(hn, "fade", 600)])
labels = []
for side, row, key, lab, col, info in PARTS:
    y = 2.2 + row * 1.42
    x = 0.6 if side == "l" else 9.03
    w = 3.7
    lw = text_width_in(lab, 15, NUNITO, True) + 0.55
    lx = x + w - lw if side == "l" else x
    pill = label(s, lx, y, lw, 0.46, lab, 15, "FFFFFF", NUNITO, True, fill=col, margins=(0.1, 0.0), name="part " + key)
    bub = glass(s, x, y + 0.55, w, 0.8, alpha=94, name="info " + key)
    shape_text(bub, info, 14, INK, FONT, True, margins=(0.14, 0.05), label="info " + key)
    px, py = PTS[key]
    tx, ty = SKX + px * SKW, SKY + py * SKH
    sx = lx + lw if side == "l" else lx
    sy = y + 0.23
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(sx), E(sy), E(tx), E(ty))
    ln.line.color.rgb = rgb(col)
    ln.line.width = Pt(3)
    lnel = ln.line._get_or_add_ln()
    lnel.set("cap", "rnd")
    lnel.append(lnel.makeelement(qn("a:tailEnd"), {"type": "oval", "w": "lg", "len": "lg"}))
    set_name(ln, "leader " + key)
    labels.append((pill, bub, ln, side))
for pill, bub, ln, side in labels:
    tl.trigger(pill, [fx(ln, "wipe_left" if side == "l" else "wipe_right", 450, sound="pop"), fx(bub, "pop", 500, delay=150)])
tl.then([fx(p, "fly_left" if sd == "l" else "fly_right", 600, delay=120 * i) for i, (p, b, l, sd) in enumerate(labels)])
next_btn(s, "Следующий зал", 10)
tl.apply()
set_transition(s, "morph", 1600, adv_click=False)
notes(s, "Зал 2. Это скелет человека. Нажимайте на цветные названия — появится линия к нужной части скелета и подсказка.\n"
         "Череп защищает головной мозг. Грудная клетка (12 пар рёбер) защищает сердце и лёгкие. Позвоночник — главная опора тела, "
         "в нём 33–34 позвонка, поэтому он гибкий. Кости таза — опора для позвоночника, к ним крепятся ноги. "
         "Кости рук помогают брать предметы, писать, рисовать (в кисти 27 косточек). Кости ног держат вес тела; бедренная кость — самая длинная.\n"
         "На этом слайде щелчок по пустому месту не перелистывает — используйте кнопку «Следующий зал» или стрелку на клавиатуре.")

# ================================================================ SLIDE 10 — hall 3 joints
s = S[10]
hall_frame(s, 3, "Суставы — «шарниры» скелета", "bg2_mint.jpg")
mn = photo(s, "mannequin", 0.6, 1.85, 4.75, 4.75, "!!hall3", focus=(0.5, 0.5),
           descr="Деревянная кукла-манекен с подвижными суставами")
dc = glass(s, 5.7, 1.85, 7.03, 1.5, name="joints def")
dtx = text(s, 6.0, 1.9, 6.5, 1.4, [[("Суставы", {"color": MINT}),
                                    (" — подвижные соединения костей. Благодаря им мы сгибаем руки и ноги, поворачиваем голову.", {})]],
           17, INK, FONT, True, name="joints def text", anchor="m")
kit.check_fit("joints", "Суставы — подвижные соединения костей. Благодаря им мы сгибаем руки и ноги, поворачиваем голову.", 17, 6.5, 1.4, FONT, True)
th = text(s, 5.7, 3.45, 7.0, 0.7, "Попробуй сам!", 36, MINT_DARK, HAND, name="try head")
ACTS = [((0.655, 0.37), 0.15, "Согни руку в локте"), ((0.74, 0.49), 0.13, "Покрути кистью"),
        ((0.39, 0.68), 0.15, "Присядь и встань"), ((0.47, 0.2), 0.19, "Наклони голову")]
tl = timeline(s)
tl.begin([fx(dc, "float_up", 500, delay=200), fx(dtx, "float_up", 500, delay=200), fx(th, "fade", 400, delay=500)])
for i, (center, side, t) in enumerate(ACTS):
    x = 5.7 + (i % 2) * 3.62
    y = 4.25 + (i // 2) * 1.2
    c = glass(s, x, y, 3.41, 1.05, name="act %d" % (i + 1))
    bd = badge(s, MANN, x + 0.14, y + 0.12, 0.81, center=center, side=side, name="act photo %d" % (i + 1))
    tt = text(s, x + 1.1, y, 2.2, 1.05, t, 16, INK, FONT, True, anchor="m", name="act text %d" % (i + 1))
    tl.click([fx(c, "pop", 500, sound="pop"), fx(bd, "spin_in", 600), fx(tt, "pop", 500)])
next_btn(s, "Следующий зал", 11)
tl.apply()
set_transition(s, "morph", 1600)
notes(s, "Зал 3. Посмотрите на деревянную куклу: она гнётся там, где у неё шарниры. У нас тоже есть такие «шарниры» — суставы.\n"
         "Суставы — подвижные соединения костей: плечевой, локтевой, коленный, тазобедренный и другие.\n"
         "По щелчку появляются задания «Попробуй сам!» — дети выполняют движения: сгибают руку в локте, крутят кистью, приседают, "
         "наклоняют голову. Спросите: какие суставы работали?")

# ================================================================ SLIDE 11 — hall 4 muscles
s = S[11]
hall_frame(s, 4, "Мышцы — двигатели тела", "bg2_peach.jpg")
ar = photo(s, "arm", 0.6, 1.85, 5.6, 4.2, "!!hall4", descr="Рука: мышцы сгибают и разгибают её в локте")
bl = label(s, 0.6 + 0.47 * 5.6 - 0.7, 1.85 + 0.25 * 4.2 - 0.2, 1.4, 0.42, "сгибает", 14, "FFFFFF", NUNITO, True,
           fill=C["coral"], margins=(0.05, 0), name="biceps label")
tlb = label(s, 0.6 + 0.3 * 5.6 - 0.8, 1.85 + 0.78 * 4.2 - 0.2, 1.6, 0.42, "разгибает", 14, "FFFFFF", NUNITO, True,
            fill=C["sky"], margins=(0.05, 0), name="triceps label")
FACTS = [(None, "мышц в теле человека — больше шестисот!"),
         ((ARM, (0.6, 0.73), 0.3), "Мышцы крепятся к костям сухожилиями. Сокращаясь, они тянут кость — и мы двигаемся."),
         ((ARM, (0.45, 0.5), 0.42), "Мышцы работают парами: одна сгибает руку, другая — разгибает."),
         ((img("mascot_flex"), (0.5, 0.3), 0.75), "Сердце — тоже мышца! Оно работает без отдыха всю жизнь.")]
tl = timeline(s)
tl.begin([fx(bl, "pop", 500, delay=400), fx(tlb, "pop", 500, delay=600)])
for i, (bdg, t) in enumerate(FACTS):
    y = 1.85 + i * 1.1
    c = glass(s, 6.55, y, 6.18, 0.96, name="fact %d" % (i + 1))
    if bdg is None:
        lead = num(s, "600+", 6.75, y, 1.6, 0.96, C["coral"], 22, anchor="m", name="fact lead 1")
        tx0 = 8.35
    else:
        path, center, side = bdg
        lead = badge(s, path, 6.72, y + 0.12, 0.72, center=center, side=side, name="fact photo %d" % (i + 1))
        tx0 = 7.6
    tt = text(s, tx0, y, 12.6 - tx0, 0.96, t, 15, INK, FONT, True, anchor="m", name="fact text %d" % (i + 1))
    tl.click([fx(c, "fly_right", 500), fx(lead, "fly_right", 500), fx(tt, "fly_right", 500)])
next_btn(s, "Следующий зал", 12)
tl.apply()
set_transition(s, "morph", 1600)
notes(s, "Зал 4. Кости сами двигаться не могут — их приводят в движение мышцы. В теле человека больше 600 мышц.\n"
         "Мышцы крепятся к костям сухожилиями. Сокращаясь, мышца тянет кость. Мышцы работают парами: "
         "двуглавая мышца плеча (бицепс) сгибает руку, трёхглавая (трицепс) — разгибает. Предложите детям согнуть руку и потрогать мышцу.\n"
         "Сердце — тоже мышца, которая работает всю жизнь без отдыха.\n"
         "Вывод: скелет и мышцы вместе — это опорно-двигательная система.")

# ================================================================ SLIDE 12 — hall 5 posture game
s = S[12]
hall_frame(s, 5, "Осанка: кто сидит правильно?", "bg2_sky.jpg")
dfn = text(s, 0.6, 1.62, 12.13, 0.5, "Осанка — это привычное положение тела, когда мы сидим, стоим и ходим.", 25,
           C["violet"], HAND, name="posture def")
IMW = 3.95
xs2 = [3.15, 3.15 + IMW + 0.85]
pics = {}
for i, key in enumerate(["posture_bad", "posture_good"]):
    x = xs2[i]
    p = photo(s, key, x, 2.3, IMW, IMW, "!!hall5" if key == "posture_good" else "posture bad", focus=(0.5, 0.55),
              descr="Мальчик сидит за партой " + ("сгорбившись" if key == "posture_bad" else "с прямой спиной"))
    lt = label(s, x - 0.25, 2.08, 0.66, 0.66, "АБ"[i], 18, "FFFFFF", DISPLAY, fill=C["violet"], oval=True,
               margins=(0, 0), name="letter %d" % (i + 1))
    pics[key] = (p, x, lt)
tl = timeline(s)
tl.begin([fx(dfn, "fade", 500), fx(pics["posture_bad"][2], "pop", 500, delay=300), fx(pics["posture_good"][2], "pop", 500, delay=450)])
m12 = mascot(s, "mascot_think", 1.55, 7.3, 2.9, name="!!mascot")
for key, good in (("posture_bad", False), ("posture_good", True)):
    p, x, _ = pics[key]
    col = C["green"] if good else C["red"]
    fr = shp(s, "round", x - 0.07, 2.23, IMW + 0.14, IMW + 0.14, radius=0.28, name="frame " + key, line=col, line_w=0.1)
    badge_t = "Правильно! Спина прямая" if good else "Неправильно! Спина согнута"
    bw = text_width_in(badge_t, 13, DISPLAY) + 0.5
    bdg = label(s, x + (IMW - bw) / 2, 2.3 + IMW - 0.38, bw, 0.6, badge_t, 13, "FFFFFF", DISPLAY, fill=col,
                margins=(0.08, 0.0), name="badge " + key)
    h = hit(s, x, 2.3, IMW, IMW, "click " + key)
    if good:
        tl.trigger(h, [fx(fr, "zoom", 400, sound="correct"), fx(bdg, "pop", 600, delay=100)]).then([fx(p, "pulse", 300, scale=105000)])
    else:
        tl.trigger(h, [fx(fr, "fade", 300, sound="wrong"), fx(bdg, "pop", 500, delay=100)]).then([fx(p, "shake")])
tl.then([fx(m12, "fly_left", 700)])
next_btn(s, "Правила осанки", 13)
tl.apply()
set_transition(s, "morph", 1600, adv_click=False)
notes(s, "Зал 5. Осанка — это привычное положение тела, когда мы сидим, стоим и ходим.\n"
         "Вопрос: кто из мальчиков сидит правильно? Пусть дети ответят, а затем нажмите на картинку.\n"
         "Картинка Б — правильно: спина прямая, ноги на полу, тетрадь далеко от глаз (зелёная рамка и звук колокольчика). "
         "Картинка А — неправильно: спина согнута, голова низко (красная рамка).\n"
         "Щелчок по пустому месту не перелистывает слайд — дальше по кнопке «Правила осанки».")

# ================================================================ SLIDE 13 — posture rules
s = S[13]
hall_frame(s, 5, "Как сохранить правильную осанку", "bg2_mint.jpg")
m13 = mascot(s, "mascot_posture", 2.15, 7.25, 5.2, name="!!mascot")
RULES = [("Сиди прямо: спина касается спинки стула, ноги стоят на полу", C["sky"]),
         ("Носи рюкзак на двух плечах, а не в одной руке", C["coral"]),
         ("Спи на ровной постели с невысокой подушкой", C["violet"]),
         ("Каждый день делай зарядку и больше двигайся", MINT)]
tl = timeline(s)
tl.begin([fx(m13, "bounce", 1000)])
for i, (t, col) in enumerate(RULES):
    x = 4.35 + (i % 2) * 4.29
    y = 1.85 + (i // 2) * 2.45
    c = glass(s, x, y, 4.09, 2.25, name="rule %d" % (i + 1))
    nn = num(s, str(i + 1), x + 0.3, y + 0.12, 1.0, 1.0, col, 44, name="rule num %d" % (i + 1))
    rh = text(s, x + 1.2, y + 0.38, 2.7, 0.5, "правило", 26, col, HAND, name="rule word %d" % (i + 1))
    tt = text(s, x + 0.3, y + 1.15, 3.55, 0.95, t, 16, INK, FONT, True, name="rule text %d" % (i + 1))
    tl.click([fx(c, "pop", 500, sound="pop"), fx(nn, "pop", 500), fx(rh, "fade", 400, delay=200), fx(tt, "fade", 400, delay=200)])
next_btn(s, "Следующий зал", 14)
tl.apply()
set_transition(s, "morph", 1600)
notes(s, "Правила хорошей осанки появляются по щелчку. Костик держит книгу на голове — попробуйте пройтись по классу так же, "
         "ровно-ровно, чтобы книга не упала! Это отличное упражнение для осанки.")

# ================================================================ SLIDE 14 — records
s = S[14]
hall_frame(s, 6, "Музей рекордов", "bg2_sun.jpg")
rc = photo(s, "records", 0.6, 1.85, 4.65, 4.65, "!!hall6", focus=(0.5, 0.5),
           descr="Лупа над крошечной косточкой и длинная бедренная кость")
RECS = [("206", "костей в скелете взрослого человека", C["violet"]),
        ("600+", "мышц помогают нам двигаться", C["coral"]),
        ("3 мм", "длина стремечка — самой маленькой косточки (она в ухе)", MINT),
        ("106", "костей в кистях и стопах — больше половины всех!", "D98A00")]
tl = timeline(s)
for i, (n_, lab, col) in enumerate(RECS):
    x = 5.6 + (i % 2) * 3.64
    y = 1.85 + (i // 2) * 2.45
    c = glass(s, x, y, 3.49, 2.2, name="record %d" % (i + 1))
    nt = num(s, n_, x + 0.1, y + 0.2, 3.29, 0.9, col, 36, align="c", name="record num %d" % (i + 1))
    lt = text(s, x + 0.25, y + 1.15, 3.0, 0.95, lab, 15, INK, FONT, True, align="c", name="record text %d" % (i + 1))
    tl.click([fx(c, "stretch_h", 450, sound="flip"), fx(nt, "pop", 600, delay=250), fx(lt, "fade", 400, delay=350)])
next_btn(s, "Выйти из музея", 15)
tl.apply()
set_transition(s, "morph", 1600)
notes(s, "Зал 6 — Музей рекордов. По щелчку открываются карточки: 206 костей у взрослого; больше 600 мышц; "
         "самая маленькая кость — стремечко в ухе, около 3 мм; в кистях и стопах 106 костей — больше половины всех костей скелета! "
         "Самая длинная кость — бедренная.\nКнопка «Выйти из музея» ведёт к физминутке.")

# ================================================================ SLIDE 15 — fizminutka
s = S[15]
bg(s, "bg2_mint.jpg")
kicker(s, "ФИЗМИНУТКА", MINT_DARK)
title(s, "Разминаем косточки и мышцы!")
nav(s)
VW = 8.3
VH = VW * 768 / 1344
mv = s.shapes.add_movie(os.path.join(ASSETS, "fizminutka.mp4"), E(0.6), E(1.85), E(VW), E(VH),
                        poster_frame_image=os.path.join(ASSETS, "fizminutka_poster.jpg"), mime_type="video/mp4")
set_name(mv, "fizminutka video")
set_descr(mv, "Видео-физминутка: учительница показывает упражнения")
mv._element.spPr.find(qn("a:prstGeom")).set("prst", "roundRect")
mv.line.color.rgb = rgb("FFFFFF")
mv.line.width = E(0.05)
shadow(mv, blur=0.35, dist=0.1, alpha=22, color="2B2A5A")
vc = text(s, 0.6, 1.85 + VH + 0.1, VW, 0.5, "Видео играет само и повторяется. Нажми на видео — пауза.", 22, MINT_DARK, HAND,
          align="c", name="video tip")
ec = glass(s, 9.25, 1.85, 3.48, 2.95, name="moves card")
eh = text(s, 9.5, 2.0, 3.0, 0.6, "Повторяй за мной!", 28, MINT_DARK, HAND, name="moves head")
el = text(s, 9.5, 2.65, 3.05, 2.1, "Руки в стороны\nНаклоны вправо и влево\nРуки вверх — потянулись!\nПрисели — встали",
          16, INK, FONT, True, para_space=5, name="moves list")
m15 = mascot(s, "mascot_flex", 11.0, 7.3, 2.5, name="!!mascot")
tl = timeline(s)
tl.begin([fx(mv, "play", 10140)])
tl.trigger(mv, [fx(mv, "toggle", 1)])
tl.media(mv, loop=True)
tl.apply()
set_transition(s, "ferris", 1500, dir="l")
notes(s, "Физминутка! Все встают. Видео запускается автоматически и повторяется по кругу — выполняйте движения вместе с героиней: "
         "руки в стороны, наклоны вправо-влево, руки вверх, присесть и встать. Щелчок по видео — пауза/продолжить.\n"
         "Видео для физминутки создано нейросетью.")

# ================================================================ SLIDE 16 — games hub
s = S[16]
bg(s, "bg2_lilac.jpg")
kicker(s, "ИНТЕРАКТИВНЫЕ ИГРЫ", C["violet"])
title(s, "Проверь себя!")
nav(s)
m16 = mascot(s, "mascot_think", 2.2, 7.25, 5.1, name="!!mascot")
GAMES = [((REC, (0.3, 0.45), 0.55), "Викторина «Знаток скелета»", "5 вопросов — выбери верный ответ", 17, C["violet"]),
         ((SKEL, (0.5, 0.29), 0.42), "Игра «Кто кого защищает?»", "Соедини кость и орган, который она бережёт", 22, C["coral"]),
         ((comp("tile_tf"), (0.5, 0.45), 0.9), "Игра «Верю — не верю»", "Переверни карточку и проверь", 23, MINT)]
tl = timeline(s)
tl.begin([fx(m16, "fly_left", 700)])
step = []
for i, ((path, center, side), nm, sub, target, col) in enumerate(GAMES):
    y = 1.9 + i * 1.62
    c = glass(s, 4.4, y, 8.33, 1.42, name="game %d" % (i + 1))
    bd = badge(s, path, 4.58, y + 0.16, 1.1, center=center, side=side, name="game photo %d" % (i + 1))
    t1 = text(s, 5.9, y + 0.2, 5.0, 0.55, nm, 17, INK, DISPLAY, name="game title %d" % (i + 1))
    t2 = text(s, 5.9, y + 0.8, 5.0, 0.45, sub, 15, C["muted"], FONT, True, name="game sub %d" % (i + 1))
    pb = label(s, 11.1, y + 0.43, 1.4, 0.56, "ИГРАТЬ", 11, "FFFFFF", DISPLAY, fill=col, margins=(0.05, 0),
               name="game go %d" % (i + 1))
    hit(s, 4.4, y, 8.33, 1.42, "game link %d" % (i + 1), target=target)
    d = 250 * i
    step += [fx(c, "fly_right", 600, delay=d), fx(bd, "fly_right", 600, delay=d), fx(t1, "fly_right", 600, delay=d),
             fx(t2, "fly_right", 600, delay=d), fx(pb, "pop", 500, delay=d + 350)]
tl.then(step)
tl.apply()
set_transition(s, "switch", 1300, dir="r")
notes(s, "Интерактивные игры, сделанные средствами PowerPoint (триггеры, анимация, гиперссылки, звуки).\n"
         "Нажмите на игру, чтобы начать. Правило: если ответ верный — зелёный цвет и колокольчик, если нет — красный цвет, "
         "и можно попробовать ещё раз. Игры идут по порядку: викторина → «Кто кого защищает?» → «Верю — не верю».")

# ================================================================ SLIDES 17–21 — quiz
QUIZ = [("Что служит опорой нашего тела?", ["Кожа", "Скелет", "Сердце"], 1),
        ("Какие кости защищают головной мозг?", ["Череп", "Рёбра", "Позвоночник"], 0),
        ("Что защищает грудная клетка?", ["Желудок", "Головной мозг", "Сердце и лёгкие"], 2),
        ("Сколько мышц в теле человека?", ["Около 60", "Больше 600", "Около 6"], 1),
        ("Как называются подвижные соединения костей?", ["Суставы", "Мышцы", "Позвонки"], 0)]
LETTERS = "АБВ"
ACOLS = [C["sky"], C["coral"], MINT]
for qi, (question, answers, right) in enumerate(QUIZ):
    s = S[17 + qi]
    bg(s, "bg2_lilac.jpg")
    kicker(s, "ВИКТОРИНА  ·  ВОПРОС %d ИЗ 5" % (qi + 1), C["violet"])
    title(s, question, size=26, w=10.6, h=1.1)
    nav(s)
    tl = timeline(s)
    m_think = mascot(s, "mascot_think", 1.5, 7.35, 2.45, name="mascot think")
    m_joy = mascot(s, "mascot_joy", 1.5, 7.35, 2.55, name="mascot joy")
    b0 = label(s, 2.85, 5.55, 4.6, 0.95, "Выбери правильный ответ!", 26, INK, HAND, fill="FFFFFF", alpha=94, radius=0.3,
               gloss=False, name="bubble ask")
    b_bad = label(s, 2.85, 5.55, 4.6, 0.95, "Ой! Подумай ещё раз", 28, C["red"], HAND, fill=C["coral_soft"], radius=0.3,
                  gloss=False, name="bubble wrong")
    b_ok = label(s, 2.85, 5.55, 4.6, 0.95, "Верно! Молодец!", 30, "138A55", HAND, fill="E3F8EC", radius=0.3, gloss=False,
                 name="bubble right")
    hits = []
    for ai, a in enumerate(answers):
        x = 0.6 + ai * 4.19
        y = 2.25
        c = glass(s, x, y, 3.75, 2.55, alpha=94, name="answer %s" % LETTERS[ai])
        lc = num(s, LETTERS[ai], x + 0.3, y + 0.15, 0.8, 0.8, ACOLS[ai], 30, name="letter %s" % LETTERS[ai])
        at = text(s, x + 0.25, y + 0.95, 3.25, 1.35, a, 19, INK, DISPLAY, align="c", anchor="m",
                  name="answer text %s" % LETTERS[ai])
        good = ai == right
        ov = label(s, x, y, 3.75, 2.55, [[("Верно!" if good else "Не то!", {"size": 26, "face": DISPLAY})],
                                         [(a, {"size": 18, "bold": True})]], 26, "FFFFFF", NUNITO,
                   fill=C["green"] if good else C["red"], radius=0.25, name="result %s" % LETTERS[ai], check=False)
        hits.append((x, y, good, ov))
    nxt_target = 18 + qi if qi < 4 else 22
    nb = next_btn(s, "Следующий вопрос" if qi < 4 else "Следующая игра", nxt_target, y=6.62)
    for x, y, good, ov in hits:
        h = hit(s, x, y, 3.75, 2.55, "choose %s" % ("right" if good else "wrong"))
        if good:
            tl.trigger(h, [fx(ov, "zoom", 450, sound="correct"), fx(b_ok, "pop", 500, delay=100),
                           fx(m_think, "fade_out", 300), fx(m_joy, "bounce", 900, delay=100)]).then([fx(nb, "fly_bottom", 500)])
        else:
            tl.trigger(h, [fx(ov, "fade", 300, sound="wrong"), fx(b_bad, "pop", 400)]
                       ).then([fx(ov, "shake")]).then([fx(b_bad, "fade_out", 400)], delay=1200)
    tl.apply()
    set_transition(s, "prism", 1400, adv_click=False)
    notes(s, "Вопрос %d. %s\nВарианты: %s.\nПравильный ответ: %s.\n"
             "Нажмите на ответ, который выбрали дети. Верный — зелёная карточка, колокольчик, Костик прыгает от радости, "
             "и появляется кнопка «%s». Неверный — красная карточка и подсказка «Подумай ещё раз»; можно выбрать снова."
          % (qi + 1, question, ", ".join(answers), answers[right], "Следующий вопрос" if qi < 4 else "Следующая игра"))

# ================================================================ SLIDE 22 — matching
s = S[22]
bg(s, "bg2_lilac.jpg")
kicker(s, "ИГРА НА СООТВЕТСТВИЕ", C["coral"])
title(s, "Кто кого защищает?")
nav(s)
ins = text(s, 0.6, 1.6, 11.5, 0.55, "Назовите пару вслух, а потом нажмите на кость — проверим!", 25, C["violet"], HAND,
           name="match hint")
LEFT = [("Череп", (0.5, 0.115), 0.24, C["violet"]), ("Грудная клетка", (0.5, 0.3), 0.36, C["coral"]),
        ("Позвоночник", (0.5, 0.41), 0.3, MINT)]
RIGHT = ["Сердце и лёгкие", "Головной мозг", "Спинной мозг"]
PAIR = {0: 1, 1: 0, 2: 2}
ROWY = [2.4, 3.9, 5.4]
tl = timeline(s)
rights = []
for j, nm in enumerate(RIGHT):
    y = ROWY[j]
    glass(s, 8.73, y, 4.0, 1.15, alpha=94, name="organ %d" % (j + 1))
    text(s, 9.05, y, 3.5, 1.15, nm, 17, INK, DISPLAY, anchor="m", name="organ text %d" % (j + 1))
    fr = shp(s, "round", 8.66, y - 0.07, 4.14, 1.29, radius=0.27, name="organ ok %d" % (j + 1), line=C["green"], line_w=0.08)
    rights.append(fr)
for i, (nm, center, side, col) in enumerate(LEFT):
    y = ROWY[i]
    glass(s, 0.6, y, 4.0, 1.15, alpha=94, name="bone card %d" % (i + 1))
    bd = badge(s, SKEL, 0.75, y + 0.13, 0.89, center=center, side=side, name="bone photo %d" % (i + 1), ring=col, ring_w=0.05)
    text(s, 1.8, y, 2.7, 1.15, nm, 17, INK, DISPLAY, anchor="m", name="bone text %d" % (i + 1))
    j = PAIR[i]
    y2 = ROWY[j]
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(4.68), E(y + 0.575), E(8.62), E(y2 + 0.575))
    ln.line.color.rgb = rgb(col)
    ln.line.width = Pt(5)
    lnel = ln.line._get_or_add_ln()
    lnel.set("cap", "rnd")
    lnel.append(lnel.makeelement(qn("a:headEnd"), {"type": "oval", "w": "sm", "len": "sm"}))
    lnel.append(lnel.makeelement(qn("a:tailEnd"), {"type": "oval", "w": "sm", "len": "sm"}))
    set_name(ln, "pair line %d" % (i + 1))
    h = hit(s, 0.6, y, 4.0, 1.15, "pick bone %d" % (i + 1))
    tl.trigger(h, [fx(ln, "wipe_left", 600, sound="correct"), fx(bd, "pulse", 250, scale=110000)]).then([fx(rights[j], "zoom", 400)])
next_btn(s, "Следующая игра", 23)
tl.apply()
set_transition(s, "conveyor", 1400, dir="l", adv_click=False)
notes(s, "Игра на соответствие. Спросите детей: что защищает череп? грудная клетка? позвоночник?\n"
         "После ответа нажмите на карточку кости — появится линия к органу, который она защищает, и зелёная рамка.\n"
         "Череп — головной мозг; грудная клетка — сердце и лёгкие; позвоночник — спинной мозг.")

# ================================================================ SLIDE 23 — true / false
s = S[23]
bg(s, "bg2_lilac.jpg")
kicker(s, "ИГРА", MINT)
title(s, "Верю — не верю")
nav(s)
ins = text(s, 0.6, 1.6, 11.5, 0.55, "Это правда или нет? Нажмите на карточку, чтобы перевернуть её", 25, C["violet"], HAND,
           name="tf hint")
TF = [("Сердце — это тоже мышца", True, "Сердце — мышца, которая работает без отдыха всю жизнь"),
      ("Позвоночник — это одна длинная кость", False, "Он состоит из 33–34 позвонков, поэтому гнётся"),
      ("Самая маленькая косточка находится в ухе", True, "Это стремечко — всего около 3 мм"),
      ("Чтобы мышцы были сильными, нужно больше лежать", False, "Мышцы крепнут, когда мы двигаемся и занимаемся спортом")]
tl = timeline(s)
for i, (stmt, truth, why) in enumerate(TF):
    x = 0.6 + i * 3.09
    y, w, h = 2.3, 2.85, 4.0
    back = label(s, x, y, w, h, [[("ВЕРЮ!" if truth else "НЕ ВЕРЮ!", {"size": 20, "face": DISPLAY})], [("", {"size": 8})],
                                 [(why, {"size": 16, "bold": True})]], 16, "FFFFFF", NUNITO,
                 fill=C["green"] if truth else C["coral"], radius=0.28, margins=(0.22, 0.2), name="tf back %d" % (i + 1))
    front = label(s, x, y, w, h, [[("?", {"size": 54, "face": DISPLAY, "color": C["lilac"]})], [(stmt, {"size": 18, "bold": True})],
                                  [("", {"size": 8})], [("верю или не верю?", {"size": 22, "face": HAND, "color": C["violet"]})]],
                  18, INK, NUNITO, fill="FFFFFF", radius=0.28, line=C["lilac"], line_w=0.03, gloss=False, margins=(0.2, 0.2),
                  name="tf front %d" % (i + 1))
    tl.trigger(front, [fx(front, "collapse_h", 260, sound="flip")]).then([fx(back, "stretch_h", 300)])
next_btn(s, "Онлайн-задание", 24)
tl.apply()
set_transition(s, "flip", 1300, dir="r", adv_click=False)
notes(s, "Игра «Верю — не верю». Учитель читает утверждение, дети поднимают руки: «верю» или «не верю». "
         "Затем нажмите на карточку — она перевернётся и покажет правильный ответ.\n"
         "1) Сердце — тоже мышца: верю. 2) Позвоночник — одна кость: не верю (33–34 позвонка). "
         "3) Самая маленькая кость в ухе: верю (стремечко, около 3 мм). 4) Чтобы мышцы были сильными, надо больше лежать: не верю.")

# ================================================================ SLIDE 24 — online task
s = S[24]
bg(s, "bg2_mint.jpg")
kicker(s, "ОНЛАЙН-ЗАДАНИЕ", C["sky"])
title(s, "Играем онлайн!")
nav(s)
m24 = mascot(s, "mascot_tablet", 2.35, 7.3, 5.4, name="!!mascot")
oc = glass(s, 4.6, 1.85, 8.13, 4.95, name="online card")
oh = text(s, 4.95, 2.05, 7.5, 0.6, "Игра «Костик и тайны скелета»", 19, INK, DISPLAY, name="online head")
LV = ["Подбери к каждой кости её работу", "«Верно или неверно» — на время", "Полезно или вредно для осанки?"]
tl = timeline(s)
tl.begin([fx(m24, "fly_left", 700), fx(oc, "float_up", 600, delay=200), fx(oh, "float_up", 600, delay=200)])
step = []
for i, t in enumerate(LV):
    y = 2.8 + i * 0.68
    nn = num(s, "0%d" % (i + 1), 4.95, y, 0.8, 0.6, [C["violet"], C["coral"], MINT][i], 18, anchor="m",
             name="level num %d" % (i + 1))
    tt = text(s, 5.8, y, 6.7, 0.6, t, 17, INK, FONT, True, anchor="m", name="level text %d" % (i + 1))
    step += [fx(nn, "fade", 400, delay=150 * i), fx(tt, "fade", 400, delay=150 * i)]
tl.then(step)
pb = label(s, 4.95, 5.05, 4.6, 0.95, "ИГРАТЬ ОНЛАЙН  »", 17, "FFFFFF", DISPLAY, fill=C["coral"], name="play online")
link_url(pb, GAME_URL)
ut = text(s, 4.95, 6.1, 5.6, 0.45, GAME_URL.replace("https://", ""), 14, C["muted"], FONT, True, name="online url")
link_url(ut, GAME_URL)
qr2 = s.shapes.add_picture(os.path.join(IMG, "qr_game.png"), E(10.8), E(4.6), E(1.65), E(1.65))
set_name(qr2, "qr game")
set_descr(qr2, "QR-код на онлайн-игру")
link_url(qr2, GAME_URL)
qt2 = text(s, 10.35, 6.22, 2.55, 0.5, "сканируй QR-код", 24, C["violet"], HAND, align="c", name="qr game text")
tl.then([fx(pb, "pop", 600, sound="pop"), fx(ut, "fade", 400), fx(qr2, "zoom", 500, delay=200), fx(qt2, "fade", 400, delay=300)])
tl.then([fx(pb, "pulse", 400, scale=107000)])
tl.apply()
set_transition(s, "airplane", 1800)
notes(s, "Онлайн-задание. Нажмите кнопку «Играть онлайн» (нужен интернет) или дайте детям отсканировать QR-код планшетом/телефоном.\n"
         "Ссылка: " + GAME_URL + "\nВ игре три уровня: подобрать к каждой кости её работу; ответить «верно/неверно» на время; "
         "разложить привычки на полезные и вредные для осанки. В конце — звёзды и диплом «Знаток скелета».")

# ================================================================ SLIDE 25 — summary: finish the sentence
s = S[25]
bg(s, "bg2_sky.jpg")
kicker(s, "ПОДВЕДЁМ ИТОГИ", C["orange"])
title(s, "Закончи предложение")
nav(s)
FIN = [("Опора нашего тела — это…", "скелет", C["violet"]), ("Головной мозг защищает…", "череп", C["sky"]),
       ("Подвижные соединения костей — это…", "суставы", MINT), ("Кости приводят в движение…", "мышцы", C["coral"]),
       ("Правильное положение тела — это…", "осанка", C["orange"])]
tl = timeline(s)
step = []
answers_ = []
for i, (q, a, col) in enumerate(FIN):
    y = 1.9 + i * 0.97
    c = glass(s, 0.6, y, 8.6, 0.82, name="sentence %d" % (i + 1))
    nn = num(s, "0%d" % (i + 1), 0.85, y, 0.8, 0.82, col, 17, anchor="m", name="sentence num %d" % (i + 1))
    qt = text(s, 1.7, y, 7.4, 0.82, q, 19, INK, FONT, True, anchor="m", name="sentence text %d" % (i + 1))
    slot = glass(s, 9.45, y + 0.07, 3.28, 0.68, alpha=55, radius=0.34, name="slot %d" % (i + 1), sh=False)
    an = label(s, 9.45, y + 0.07, 3.28, 0.68, a.upper(), 17, "FFFFFF", DISPLAY, fill=col, margins=(0.1, 0),
               name="answer word %d" % (i + 1))
    d = 120 * i
    step += [fx(c, "fly_left", 500, delay=d), fx(nn, "fly_left", 500, delay=d), fx(qt, "fly_left", 500, delay=d),
             fx(slot, "fade", 400, delay=d + 200)]
    answers_.append(an)
tl.begin(step)
for an in answers_:
    tl.click([fx(an, "pop", 600, sound="correct")])
tl.apply()
set_transition(s, "vortex", 1800, dir="r")
notes(s, "Подведение итогов в игровой форме. Учитель читает начало фразы, дети хором заканчивают, затем щелчок — появляется ответ.\n"
         "Ответы: скелет; череп; суставы; мышцы; осанка.")

# ================================================================ SLIDE 26 — infographic
s = S[26]
bg(s, "bg2_sun.jpg")
kicker(s, "ИНФОГРАФИКА", "B36B00")
title(s, "Что мы узнали")
nav(s)
cc = glass(s, 4.87, 2.0, 3.6, 3.6, alpha=70, radius=1.8, name="center circle")
mj = mascot(s, "mascot_joy", 6.67, 5.55, 3.3, name="!!mascot")
fp = label(s, 4.37, 5.85, 4.6, 0.75, "Скелет + мышцы = движение!", 14, "FFFFFF", DISPLAY, fill=C["indigo"],
           margins=(0.1, 0), name="formula")
INFO = [(0.6, 1.95, (SKEL, (0.5, 0.3), 0.6), "Скелет", "206 костей. Опора и защита тела", C["violet"]),
        (8.83, 1.95, (MANN, (0.665, 0.39), 0.25), "Суставы", "Подвижные соединения: сгибаемся и поворачиваемся", MINT),
        (0.6, 4.3, (ARM, (0.47, 0.45), 0.5), "Мышцы", "Больше 600. Приводят кости в движение", C["coral"]),
        (8.83, 4.3, (img("posture_good"), (0.62, 0.4), 0.62), "Осанка", "Прямая спина — здоровое тело", C["sky"])]
tl = timeline(s)
tl.begin([fx(cc, "zoom", 700), fx(mj, "bounce", 1000, delay=200)])
step = []
for i, (x, y, (path, center, side), hd, tx, col) in enumerate(INFO):
    c = glass(s, x, y, 3.9, 2.05, name="info card %d" % (i + 1))
    bd = badge(s, path, x + 0.22, y + 0.22, 0.95, center=center, side=side, name="info photo %d" % (i + 1), ring=col)
    h1 = text(s, x + 1.35, y + 0.35, 2.45, 0.6, hd, 19, col, DISPLAY, name="info head %d" % (i + 1))
    t1 = text(s, x + 0.25, y + 1.25, 3.45, 0.75, tx, 15, INK, FONT, True, name="info text %d" % (i + 1))
    step += [fx(c, "pop", 600, delay=250 * i), fx(bd, "spin_in", 700, delay=250 * i), fx(h1, "fade", 400, delay=250 * i + 200),
             fx(t1, "fade", 400, delay=250 * i + 200)]
tl.then(step)
tl.then([fx(fp, "pop", 600, sound="win")])
tl.apply()
set_transition(s, "honeycomb", 1800)
notes(s, "Инфографика «Что мы узнали»: скелет — 206 костей, опора и защита тела; суставы — подвижные соединения костей; "
         "мышцы — больше 600, приводят кости в движение; осанка — правильное положение тела. "
         "Скелет и мышцы вместе — опорно-двигательная система: «Скелет + мышцы = движение!»")

# ================================================================ SLIDE 27 — keep healthy
s = S[27]
bg(s, "bg2_mint.jpg")
kicker(s, "СОВЕТЫ КОСТИКА", MINT_DARK)
title(s, "Крепкие кости и сильные мышцы")
nav(s)
fd = photo(s, "food", 0.6, 1.85, 5.0, 3.75, "food", focus=(0.5, 0.62), descr="Молоко, сыр, творог, рыба и другие полезные продукты")
sp = photo(s, "sport", 5.85, 1.85, 6.88, 3.75, "sport", focus=(0.5, 0.5), descr="Дети занимаются спортом в парке")
f1 = glass(s, 0.6, 5.8, 5.0, 1.05, name="food tip")
ft = text(s, 0.85, 5.8, 4.6, 1.05, "Ешь молоко, творог, сыр и рыбу — в них кальций для костей", 15, INK, FONT, True,
          anchor="m", name="food tip text")
f2 = glass(s, 5.85, 5.8, 6.88, 1.05, name="sport tip")
st = text(s, 6.1, 5.8, 6.4, 1.05, "Двигайся: зарядка, прогулки, спорт. И высыпайся — 9–10 часов!", 15, INK, FONT, True,
          anchor="m", name="sport tip text")
tl = timeline(s)
tl.begin([fx(fd, "zoom", 700), fx(sp, "zoom", 700, delay=200)])
tl.click([fx(f1, "float_up", 500), fx(ft, "float_up", 500)])
tl.click([fx(f2, "float_up", 500), fx(st, "float_up", 500)])
tl.apply()
set_transition(s, "pan", 1300, dir="u")
notes(s, "Как сделать кости крепкими, а мышцы сильными? Правильно питаться: молоко, творог, сыр, рыба содержат кальций. "
         "Много двигаться: зарядка, прогулки, спорт. Высыпаться: младшим школьникам нужно 9–10 часов сна. И следить за осанкой!")

# ================================================================ SLIDE 28 — reflection ladder
s = S[28]
bg(s, "bg2_hero_end.jpg")
kicker(s, "РЕФЛЕКСИЯ", C["sun"])
title(s, "Лесенка успеха", color="FFFFFF")
nav(s)
sub = text(s, 0.6, 1.6, 5.2, 0.6, "Нажми на ступеньку, где ты сейчас", 27, "E9E4FF", HAND, name="ladder hint")
STEPS = [(1.0, 4.6, 2.45, C["coral"], "Мне было трудно", "Не страшно! Повтори дома — и всё получится", "mascot_shrug", "pop"),
         (4.87, 3.7, 3.35, C["sun"], "Есть вопросы", "Хорошо! Спроси учителя — и станешь знатоком", "mascot_think", "pop"),
         (8.74, 2.8, 4.25, "1FB89A", "Всё понятно! Могу объяснить другу", "Супер! Ты настоящий знаток скелета!", "mascot_joy", "win")]
tl = timeline(s)
for i, (x, y, h, col, lab, msg, mkey, snd) in enumerate(STEPS):
    w = 3.6
    tc = INK if col == C["sun"] else "FFFFFF"
    st_ = label(s, x, y, w, h, lab, 15, tc, DISPLAY, fill=col, radius=0.22, anchor="t", margins=(0.2, 0.2),
                name="step %d" % (i + 1))
    ms = text(s, x + 0.2, y + 1.0, w - 0.4, 1.0, msg, 24, tc, HAND, name="step msg %d" % (i + 1), spacing=0.9)
    mm = mascot(s, mkey, x + w / 2, y + 0.05, 2.05, name="step mascot %d" % (i + 1))
    tl.trigger(st_, [fx(mm, "bounce", 900, sound=snd), fx(ms, "fade", 500, delay=300)])
tl.apply()
set_transition(s, "glitter", 1600, dir="l", pattern="hexagon")
notes(s, "Рефлексия «Лесенка успеха». Каждый ученик решает, на какой ступеньке он сейчас: "
         "«Мне было трудно», «Есть вопросы» или «Всё понятно! Могу объяснить другу». "
         "Нажмите на ступеньку — на неё запрыгнет Костик и скажет ободряющие слова. Можно попросить детей поднять руку у своей ступеньки.")

# ================================================================ SLIDE 29 — finish the phrase
s = S[29]
bg(s, "bg2_lilac.jpg")
kicker(s, "РЕФЛЕКСИЯ", C["violet"])
title(s, "Продолжи фразу")
nav(s)
PH = [((SKEL, (0.5, 0.12), 0.3), "Сегодня на уроке я узнал(а)…", C["sky"]),
      ((comp("tile_reflect"), (0.5, 0.42), 0.9), "Больше всего мне понравилось…", C["coral"]),
      ((img("posture_good"), (0.62, 0.42), 0.6), "Я буду беречь свою осанку, потому что…", MINT)]
tl = timeline(s)
step = []
cards = []
for i, ((path, center, side), t, col) in enumerate(PH):
    x = 0.6 + i * 4.115
    c = glass(s, x, 2.0, 3.9, 3.3, alpha=94, radius=0.35, name="phrase %d" % (i + 1))
    bd = badge(s, path, x + 0.3, 2.3, 1.15, center=center, side=side, name="phrase photo %d" % (i + 1), ring=col)
    tt = text(s, x + 0.3, 3.65, 3.3, 1.5, t, 30, INK, HAND, name="phrase text %d" % (i + 1), spacing=0.95)
    star = badge(s, REC, x + 2.75, 1.7, 1.0, center=(0.875, 0.77), side=0.2, name="phrase star %d" % (i + 1), ring=C["sun"])
    step += [fx(c, "pop", 600, delay=200 * i), fx(bd, "pop", 600, delay=200 * i), fx(tt, "fade", 400, delay=200 * i + 250)]
    cards.append((c, star))
bot = text(s, 0.6, 5.55, 12.13, 0.6, "Ответил — нажми на карточку и получи звезду от Костика!", 28, C["indigo"], HAND,
           align="c", name="phrase hint")
tl.begin(step)
tl.then([fx(bot, "fade", 500)])
for c, star in cards:
    tl.trigger(c, [fx(star, "spin_in", 700, sound="correct")])
tl.apply()
set_transition(s, "pageCurlDouble", 1600)
notes(s, "Продолжите фразу. Дети по желанию заканчивают одну из фраз. После ответа нажмите на карточку — Костик дарит звезду.")

# ================================================================ SLIDE 30 — thanks
s = S[30]
b30 = bg(s, "bg2_hero_end.jpg")
ttl = big_title(s, [("Спасибо", "FFFFFF"), ("за урок!", C["sun"])], y=1.55)
sub = text(s, 0.8, 4.05, 7.2, 0.9, "Берегите осанку и будьте здоровы!", 34, "E9E4FF", HAND, name="thanks sub")
au = text(s, 0.8, 5.45, 6.0, 0.9, [[("Шелешко А.И.", {"size": 19, "bold": True, "color": "FFFFFF"})],
                                   [("группа П-31", {"size": 15, "color": "D9D2FF", "bold": True})]], name="thanks author")
mt = mascot(s, "mascot_wave", 10.45, 7.25, 6.1, name="!!mascot")
tl = timeline(s)
tl.begin([fx(b30, "grow", 22000, scale=108000), fx(ttl, "pop", 900, sound="win"), fx(mt, "bounce", 1200, delay=200)])
tl.then([fx(sub, "float_up", 600)])
tl.then([fx(au, "fade", 500)])
tl.then([fx(mt, "bob", 1700, dy="-0.018")])
tl.apply()
set_transition(s, "prestige", 1800)
notes(s, "Спасибо за урок! Костик прощается с ребятами. Напомните: берегите осанку, больше двигайтесь и правильно питайтесь.")

# ================================================================ SLIDE 31 — sources & tools
s = S[31]
bg(s, "bg2_sky.jpg")
kicker(s, "ИСТОЧНИКИ И ИНСТРУМЕНТЫ", C["indigo"])
title(s, "Над уроком работали")
nav(s)
glass(s, 0.6, 1.85, 5.6, 4.95, name="sources card")
text(s, 0.9, 2.05, 5.0, 0.5, "Источники", 18, INK, DISPLAY, name="sources head")
text(s, 0.9, 2.65, 5.1, 1.0, "Плешаков А.А. Окружающий мир. 3 класс. Учебник, часть 1. Тема «Опора тела и движение»", 15,
     INK, FONT, False, name="source book")
src2 = text(s, 0.9, 3.75, 5.1, 0.5, [[("Видеоэкскурсия: ", {"bold": True}), ("youtu.be/6C_RWeZtCNg", {"color": C["violet"], "bold": True})]],
            15, name="source video")
link_url(src2, VIDEO_URL)
src3 = text(s, 0.9, 4.3, 5.1, 0.5, [[("Онлайн-игра: ", {"bold": True}), (GAME_URL.replace("https://", ""), {"color": C["violet"], "bold": True})]],
            15, name="source game")
link_url(src3, GAME_URL)
text(s, 0.9, 6.0, 5.1, 0.6, "Шрифты Unbounded, Nunito и Caveat (лицензия SIL OFL)", 13, C["muted"], FONT, True, name="source fonts")
text(s, 6.55, 1.85, 6.0, 0.5, "Нейросети в проекте", 18, INK, DISPLAY, name="ai head")
AI = [("Claude (Anthropic)", "сценарий урока, тексты, вопросы, вёрстка презентации, онлайн-игра", C["violet"],
       (img("mascot_think"), (0.5, 0.3), 0.8)),
      ("Nano Banana 2 (Google) в Higgsfield", "иллюстрации и маскот Костик", C["coral"], (img("mascot_wave"), (0.5, 0.3), 0.8)),
      ("MiniMax Hailuo", "видео для физминутки", MINT, (os.path.join(ASSETS, "fizminutka_poster.jpg"), (0.45, 0.35), 0.75))]
for i, (nm, what, col, (path, center, side)) in enumerate(AI):
    y = 2.5 + i * 1.45
    glass(s, 6.55, y, 6.18, 1.28, name="ai %d" % (i + 1))
    badge(s, path, 6.72, y + 0.17, 0.94, center=center, side=side, name="ai photo %d" % (i + 1), ring=col)
    text(s, 7.85, y + 0.14, 4.75, 0.45, nm, 13, INK, DISPLAY, name="ai name %d" % (i + 1))
    text(s, 7.85, y + 0.58, 4.75, 0.65, what, 14, C["muted"], FONT, True, name="ai what %d" % (i + 1))
set_transition(s, "fade", 900)
notes(s, "Источники и инструменты. Нейросети, которые использовались при подготовке проекта: Claude — сценарий, тексты, задания, "
         "вёрстка и онлайн-игра; Nano Banana 2 (через Higgsfield) — иллюстрации; MiniMax Hailuo — видео для физминутки.")

prs.save(OUT)
print("saved", OUT)
print("fit warnings:")
for w in kit.WARNINGS:
    print("  ", w)
