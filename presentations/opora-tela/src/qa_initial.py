"""Copy of a deck with every shape removed whose first animation is an entrance —
i.e. what the audience sees right after the slide appears."""
import sys
from pptx import Presentation
from lxml import etree
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
prs = Presentation(sys.argv[1])
for slide in prs.slides:
    sld = slide._element
    first = {}
    for ctn in sld.iter(P + "cTn"):
        cls = ctn.get("presetClass")
        if not cls:
            continue
        for tgt in ctn.iter(P + "spTgt"):
            first.setdefault(tgt.get("spid"), cls)
            break
    for shp in list(slide.shapes):
        if first.get(str(shp.shape_id)) == "entr":
            shp._element.getparent().remove(shp._element)
prs.save(sys.argv[2])
